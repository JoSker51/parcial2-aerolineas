"""Genera docs/PARCIAL2-DOCUMENTO.docx (documento de entrega del parcial 2).

Uso: python docs/build_doc.py
El índice (TOC) lo rellena Word al abrir el archivo, o docs/word_update.ps1 (actualiza y exporta PDF).
"""
import pathlib
import re

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
FIG = HERE / "diagramas"
OUT = HERE / "PARCIAL2-DOCUMENTO.docx"

BLUE = RGBColor(0x1F, 0x4F, 0xC4)
INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x5B, 0x64, 0x75)
HEAD_FILL = "1F4FC4"
ZEBRA = "F2F5FB"

doc = Document()

# ------------------------------------------------------------------ estilos base
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)          # carta
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.3))
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)

st = doc.styles["Normal"]
st.font.name, st.font.size, st.font.color.rgb = "Calibri", Pt(11), INK
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12
for lvl, size in ((1, 17), (2, 13.5), (3, 12)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = BLUE if lvl < 3 else INK
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.paragraph_format.space_before = Pt(16 if lvl == 1 else 11)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True

# pie de página con número
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _field(par, instr):
    r = par.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
    return r


fr = fp.add_run("Parcial 2 · Big Data e Ingeniería de Datos · pág. ")
fr.font.size, fr.font.color.rgb = Pt(9), MUTED
_field(fp, "PAGE").font.size = Pt(9)

# ------------------------------------------------------------------ helpers
TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[\[.+?\]\])")


def runs(par, text, size=None, color=None):
    """Texto con **negrita**, `código` y [[resaltado para completar]]."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 11) - 1.5)
            r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
        elif part.startswith("[["):
            r = par.add_run(part[2:-2]); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = par.add_run(part)
        if size and not part.startswith("`"):
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return par


def P(text, size=None, color=None, align=None, after=None):
    p = runs(doc.add_paragraph(), text, size, color)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def H(text, lvl=1):
    return doc.add_heading(text, lvl)


def B(items, style="List Bullet"):
    for it in items:
        p = runs(doc.add_paragraph(style=style), it)
        p.paragraph_format.space_after = Pt(3)


def N(items):
    B(items, "List Number")


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def T(header, rows, widths, size=9.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        _shade(c, HEAD_FILL)
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            runs(cells[i].paragraphs[0], str(v), size)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            if k % 2 == 1:
                _shade(cells[i], ZEBRA)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # repetir encabezado al cambiar de página
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def IMG(path, width_cm, caption):
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = P(caption, 9, MUTED, WD_ALIGN_PARAGRAPH.CENTER)
    c.runs[0].italic = True


def NOTE(text):
    """Párrafo destacado con borde izquierdo (decisión o conclusión clave)."""
    p = runs(doc.add_paragraph(), text, 10.5)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "24"), ("space", "8"), ("color", "1F4FC4")):
        left.set(qn(f"w:{k}"), v)
    bdr.append(left); pPr.append(bdr)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "EEF4FF"); pPr.append(shd)
    p.paragraph_format.left_indent = Cm(0.3)
    return p


def PAGEBREAK():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================== PORTADA
for _ in range(5):
    doc.add_paragraph()
P("SEGUNDO PARCIAL", 12, MUTED, WD_ALIGN_PARAGRAPH.CENTER, 2)
P("**Sistema de reservas de aerolínea**", 26, BLUE, WD_ALIGN_PARAGRAPH.CENTER, 0)
P("Panel analítico gerencial con EMR, Glue Data Catalog y Athena", 15, INK, WD_ALIGN_PARAGRAPH.CENTER, 24)
P("Big Data e Ingeniería de Datos", 12, MUTED, WD_ALIGN_PARAGRAPH.CENTER, 40)
P("Integrantes: [[Nombre completo integrante 1]] · [[Nombre completo integrante 2]]", 11.5, None, WD_ALIGN_PARAGRAPH.CENTER, 4)
P("Docente: [[Nombre del profesor]]", 11.5, None, WD_ALIGN_PARAGRAPH.CENTER, 4)
P("Seed del equipo: **190** · Cuenta AWS Academy Learner Lab, región us-east-1", 11.5, None, WD_ALIGN_PARAGRAPH.CENTER, 4)
P("Entrega: 23 de octubre de 2026", 11.5, None, WD_ALIGN_PARAGRAPH.CENTER, 40)

T(["Resultado del harness (19-oct-2026)", "Costo mensual proyectado", "Costo estimado por el harness (en vivo)"],
  [["**70,0 / 70,0** — los 12 requisitos pasan", "**US$ 19,83** (límite US$ 50)", "US$ 15,90"]],
  [5.6, 5.4, 5.6], 10.5)
PAGEBREAK()

# ================================================================== ÍNDICE
P("**Contenido**", 15, BLUE)
_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
PAGEBREAK()

# ================================================================== 1. RESUMEN
H("1. Resumen")
P("El primer parcial diseñó el sistema de reservas y construyó su backend. En este parcial el sistema funciona de punta a punta y se mide con el harness del docente. Tiene **dos caminos que no se tocan**:")
B(["**Camino transaccional (clientes):** una sola EC2 `t3.small` corre la API FastAPI, sirve las dos páginas web y aloja PostgreSQL 15. Aquí viven la búsqueda (R1), la reserva sin sobreventa (R2), la consulta de reservas (R3) y el flujo web del cliente (R8).",
   "**Camino analítico (gerentes):** el histórico sucio (2,07 M filas en CSV) se sube a S3. Un cluster EMR efímero lo limpia con PySpark (reglas C1–C5) y escribe Parquet particionado por mes. Las tablas se registran en el Glue Data Catalog y Athena las consulta. El endpoint `/analytics/occupancy` y el panel `/manager` leen de Athena (R4–R7, R9)."])
P("La analítica **nunca lee PostgreSQL**. Por eso un gerente consultando el histórico no puede frenar una reserva, que es la exigencia de separación de cargas del enunciado.")
NOTE("**Resultado verificado:** el harness del docente (`pytest`) obtuvo **70/70** contra el sistema desplegado: p95 de búsqueda 0,185 s; 0 sobreventas en 10 rondas de 50 solicitudes simultáneas; 2.005.463 filas limpias exactas; top-10 y ocupaciones idénticas al oráculo; endpoint de Athena en 1,6–2,2 s; 24 horas normalizadas de EMR usadas de 150; proyección de costos coherente con lo desplegado. El reporte completo está en la sección 7 y en `report.md`.")

# ================================================================== 2. ARQUITECTURA
H("2. Arquitectura desplegada")
IMG(FIG / "arquitectura.png", 16.9, "Figura 1. Arquitectura desplegada. Arriba, el camino transaccional; abajo, el camino analítico. Cada recuadro naranja es un recurso real de la cuenta AWS.")
T(["Componente", "Recurso real en AWS", "Función"],
  [["Cómputo transaccional", "EC2 `t3.small`, Amazon Linux 2023, Elastic IP 100.60.245.136, perfil `LabInstanceProfile`", "API (puerto 8000), páginas `/` y `/manager`, PostgreSQL local"],
   ["Base transaccional", "PostgreSQL 15 en la misma EC2 (solo `localhost`), disco gp3 de 8 GB", "Tablas `flight` y `reservation`; CHECK anti-sobreventa"],
   ["Data lake", "S3 `airline-p2-488483603895` (acceso público bloqueado)", "`raw/` (CSV), `curated/` (Parquet), resultados de Athena y logs de EMR con expiración"],
   ["Procesamiento batch", "EMR 6.15 (Spark 3.4), 1 master + 1 core `m4.large`, con auto-terminación", "`emr/curate.py`: reglas C1–C5, cruce con vuelos y métricas"],
   ["Catálogo", "Glue Data Catalog, base `airline_analytics`", "`reservations_curated` (projection por mes) y `route_month_metrics`"],
   ["Consulta analítica", "Athena, workgroup `primary`", "Consultas del harness, del endpoint y del panel"]],
  [3.6, 6.6, 6.6])

# ================================================================== 3. DECISIONES (SECCIÓN 5)
H("3. Decisiones de diseño y su justificación (sección 5)")
P("Cada decisión cita el requisito que la motiva y el pilar del AWS Well-Architected Framework que la sostiene: Excelencia Operacional (EO), Seguridad (SEG), Fiabilidad (FIA), Eficiencia de Rendimiento (REN), Optimización de Costos (COS) y Sostenibilidad (SOS).", 10.5, MUTED)

H("3.1 Modelo de hosting: una sola EC2 con PostgreSQL local", 2)
P("**Decisión.** Backend, frontend y base transaccional corren en **una EC2 `t3.small`**. FastAPI sirve la API y también las dos páginas estáticas, con un solo proceso, un solo puerto y sin CORS. PostgreSQL 15 está instalado en la misma máquina y solo escucha en `localhost`.")
P("**Por qué es la opción más barata que cumple R1 y R2.** R1 pide p95 ≤ 1 s con 20 clientes y R2 pide serializar 50 reservas simultáneas sin errores 5xx. Las dos cosas las resuelve un solo motor PostgreSQL con bloqueo de fila; no hace falta más infraestructura. Medido: p95 = 0,185 s y 0 sobreventas.")
T(["Alternativa", "Costo/mes", "Por qué no"],
  [["**EC2 t3.small + PostgreSQL local (elegida)**", "**19,47**", "Cumple R1 y R2 con margen; un servicio menos que con RDS"],
   ["EC2 t3.micro + RDS db.t3.micro (20 GB)", "26,59", "+37 %. Con 1 GB de RAM la EC2 queda justa para 2 procesos de uvicorn"],
   ["EC2 t3.small + RDS db.t3.micro", "34,18", "+76 %. RDS aporta backups automáticos, pero ningún requisito los exige"],
   ["ECS Fargate + RDS + ALB", "> 45", "El ALB solo ya cuesta ~16 USD/mes; el harness lo marca como advertencia en R11"],
   ["Lambda + API Gateway + DynamoDB", "variable", "Reescribir el control de sobreventa sobre otro motor; API Gateway suma un servicio más (R11)"]],
  [6.5, 2.2, 8.1])
P("**Más o menos administrado.** RDS sería *más* administrado: backups, parches y failover. Se descartó porque el enunciado premia lo más barato que cumple y porque en este laboratorio RDS no puede ser Multi-AZ (R11), así que su ventaja de fiabilidad sería parcial. Lo que se paga se declara en la sección 8: los backups quedan a cargo del equipo. **Pilares:** COS (−US$ 12/mes), REN (la consulta no cruza la red hasta otro servidor), FIA (compromiso declarado).")
P("**Por qué `t3.small` y no `t3.micro`.** La instancia aloja PostgreSQL, 2 procesos de uvicorn y el agente de SSM. Con 1 GB (micro) quedaría expuesta a quedarse sin memoria durante la ráfaga de R2. La diferencia es de US$ 7,6/mes.")
P("**IP elástica.** El Learner Lab detiene la EC2 al cerrar la sesión, y al reiniciarla cambia la IP pública. El docente califica con las URLs del `config.yaml` entregado, así que una IP fija evita entregar una URL muerta. Cuesta lo mismo que la IP pública automática (US$ 0,005/h) y no agrega servicios. **Pilares:** FIA, EO.")

H("3.2 Por qué EMR, qué hace y cómo se garantiza que se apague", 2)
P("**Qué hace el job que Athena sola no podría.** `emr/curate.py` aplica las cinco reglas en el mismo orden que el oráculo y **materializa una sola vez** datos limpios, tipados y particionados:")
T(["Regla", "Implementación en PySpark", "Detalle que importa"],
  [["C1 ruta", "`upper(regexp_replace(route, '\\s+', ''))` sobre **flights**", "La ruta sucia trae espacios en el medio (`BOG - MDE`): un `trim` no basta. Las reservas crudas no traen ruta: sale del vuelo"],
   ["C2 status", "`upper(regexp_replace(status, '^\\s+|\\s+$', ''))`", "Equivale al `strip()` de Python (cualquier espacio en los extremos, no solo ' ')"],
   ["C3", "`status IN ('CONFIRMED','CANCELLED')`", "Descarta PENDING y otros estados"],
   ["C4", "`price_cents IS NOT NULL AND price_cents > 0`", "Con esquema explícito, el precio vacío del CSV se lee como NULL"],
   ["C5", "`row_number()` por `reservation_id`, ordenado por la posición de la fila en el archivo", "Conserva la **primera** aparición. `dropDuplicates` se quedaría con una cualquiera"]],
  [1.7, 6.6, 8.5])
P("Athena sola tendría que limpiar 85 MB de CSV **en cada consulta**. Además, no tiene forma confiable de saber cuál fila apareció primero en el archivo, que es lo que exige C5.")
T(["Alternativa considerada", "Por qué no"],
  [["Athena CTAS", "Puede convertir a Parquet y particionar, pero no garantiza la \"primera aparición\" de C5. Tampoco cumple R4, que exige que EMR produzca los datos"],
   ["AWS Glue job (Spark)", "Usaría el mismo código y costaría parecido o menos (~US$ 0,05 por corrida, 2 DPU durante unos minutos). Se descarta porque el curso exige EMR y porque R4 verifica corridas de EMR"],
   ["EMR Serverless", "Elimina el aprovisionamiento, pero hay que dejar la aplicación detenida (R10) y su disponibilidad en el Learner Lab es incierta. El cluster con auto-terminación es más simple de garantizar"]],
  [4.2, 12.6])
P("**Opción e instancias.** EMR sobre EC2 (release 6.15, Spark 3.4) con **1 master + 1 core `m4.large`**, el único tipo permitido en el laboratorio. El job corre como *step* en modo `client` con executors de 2 GB. El cluster vive unos 9 minutos (aprovisionar, configurar, ~2 min de job, apagar). Cada corrida factura 8 horas normalizadas: con 24 ya usadas, el tope de R10 (150) deja margen para ~15 corridas más.")
NOTE("**Se probó primero con un solo nodo y falló dos veces** (hallazgos H-1 y H-2, sección 7.3). En la primera, el driver repetía `Initial job has not accepted any resources`: con la memoria por defecto de EMR, el driver y un executor no caben en los ~6 GB de YARN de un m4.large. En la segunda, el nodo único quedó saturado alojando HDFS, YARN, el driver y los executors, y ni siquiera subía sus logs a S3. Con un nodo core separado el job tarda 2 minutos. Por eso el enunciado pide probar con una muestra: ambos fallos se detectaron con 150 mil filas, no con 2 millones.")
P("**Cómo se garantiza el apagado (R10).** Tres mecanismos:")
B(["`--auto-terminate`: el cluster se destruye al terminar el último step, sin intervención humana.",
   "`ActionOnFailure=TERMINATE_CLUSTER`: si el job falla, el cluster también se apaga.",
   "El script espera `cluster-terminated` y reporta las horas normalizadas usadas."])
P("Evidencia del harness: `active clusters=[]`, 24 horas normalizadas usadas de 150. **Pilares:** COS y SOS (no hay cómputo ocioso), EO (el apagado no depende de acordarse).")

H("3.3 Disposición de los datos: Parquet particionado por mes", 2)
P("**Por qué Parquet.** Es columnar y comprimido: el CSV de 85 MB queda en **~20 MB** de Parquet (4×), con tipos reales (`bigint`, `double`). Athena lee solo las columnas que pide la consulta y cobra por byte leído.")
P("**Por qué particionar por mes.** Todas las preguntas del negocio filtran por mes (\"cómo le fue a una ruta en un mes\"). Con `month=YYYY-MM/` como carpeta, un `WHERE month = '2024-01'` lee solo 1 de 24 carpetas. Se usa *partition projection*: Athena deduce las 24 particiones de las propiedades de la tabla, sin `MSCK REPAIR` ni registrar particiones después de cada corrida.")
T(["Medición (harness R6)", "Bytes", "Comentario"],
  [["Tamaño promedio de una partición", "829.146", "24 particiones, ~19,9 MB en total"],
   ["Consulta de un mes (`count`, `sum(price_cents)`)", "**295.273**", "Lee una partición y solo la columna de precio"],
   ["Límite de R6 (2 × partición promedio)", "1.658.293", "Holgura de 5,6×"],
   ["Misma consulta sin particionar (estimado)", "~7.100.000", "Leería la columna de precio de los 24 meses: 24× más"],
   ["Misma consulta sobre el CSV crudo", "~84.400.000", "Leería el archivo completo: ~285× más"]],
  [7.2, 2.8, 6.8])
P("**Matiz honesto.** Athena factura un mínimo de 10 MB por consulta, así que a este volumen la partición casi no cambia la factura. Sí cambia los **bytes leídos**, que es lo que mide R6, y escala bien: con 10× más datos, el CSV costaría 10× más por consulta y la partición mantendría cada consulta acotada a un mes. **Pilares:** REN, COS, SOS (menos datos almacenados y leídos).")

H("3.4 Cómo obtiene el panel sus números", 2)
P("`GET /analytics/occupancy` ejecuta en Athena, en cada petición, una consulta **parametrizada** sobre `route_month_metrics`. La ruta y el mes viajan como `ExecutionParameters`, no concatenados en el SQL, así que no hay inyección SQL posible. El panel `/manager` llama a ese mismo endpoint y muestra el valor con dos decimales.")
T(["Aspecto", "Consecuencia de consultar Athena en cada petición"],
  [["Latencia", "1,6–2,2 s medidos (R7 exige ≤ 10 s). Aceptable para un gerente; inaceptable para el flujo de compra, que por eso nunca toca Athena"],
   ["Costo por consulta", "La tabla de métricas tiene 336 filas (unos KB), pero se factura el mínimo de 10 MB: ~US$ 0,00005 por consulta. 1.000 consultas/mes ≈ US$ 0,05"],
   ["Frescura", "Los datos son tan frescos como la última corrida de EMR (mensual). El panel responde preguntas históricas y tácticas, no la venta del minuto"],
   ["¿Caché?", "**No se implementó.** Ahorraría como máximo ~US$ 0,40/mes incluso con 10× más consultas, y obligaría a invalidar después de cada corrida de EMR. Se justificaría con latencias objetivo < 1 s o con miles de consultas por minuto. La opción más barata sería activar *query result reuse* de Athena, sin código. El contrato permite caché solo si la primera petición de cada (ruta, mes) consulta Athena"]],
  [3.0, 13.8])
P("**Pilares:** COS (no hay base analítica encendida: el parcial 1 tenía un RDS analítico de US$ 15,44/mes), SEG (consulta parametrizada; credenciales del rol de instancia, no claves en disco), REN.")

H("3.5 Separación de cargas: la analítica nunca frena las reservas", 2)
B(["**Cómputo separado.** La limpieza corre en nodos EMR propios que existen solo durante el job. Las consultas corren en Athena, que es serverless. Ninguna de las dos usa la CPU ni la memoria de la EC2 transaccional.",
   "**Datos separados.** La analítica lee S3, nunca PostgreSQL. El histórico llega como CSV exportado y la base transaccional solo atiende búsquedas y reservas.",
   "**El único punto compartido es el proceso web.** Una petición a `/analytics/occupancy` espera la respuesta de Athena, pero esa espera es de red: FastAPI corre los endpoints síncronos en un pool de hilos y la petición no bloquea las reservas. El pool de conexiones a PostgreSQL ni siquiera se usa en esa ruta.",
   "**El bloqueo se mantiene corto.** La reserva bloquea **una fila** (`SELECT … FOR UPDATE`) durante una transacción de milisegundos, sin llamadas externas dentro (DEC-4 del parcial 1). La búsqueda no toma bloqueos."])
P("Evidencia: R1 y R2 pasaron con el camino analítico desplegado y en uso. **Pilares:** REN, FIA.")

H("3.6 Tabla de justificación de servicios", 2)
T(["Servicio", "Alternativa considerada", "Requisito", "Pilar(es)"],
  [["**Amazon EC2** (`t3.small` + Elastic IP)", "ECS Fargate + ALB; Lambda + API Gateway; Elastic Beanstalk", "R1, R2, R3, R8, R11, R12", "COS, EO"],
   ["**PostgreSQL 15** (en la EC2)", "RDS db.t3.micro; DynamoDB con escritura condicional", "R2, R3, R11, R12", "COS, REN; FIA como compromiso"],
   ["**Amazon S3**", "HDFS del cluster (se pierde al apagarlo); EBS", "R4, R5, R6", "FIA (11 nueves), COS, SOS"],
   ["**Amazon EMR** (cluster efímero `m4.large`)", "Glue job Spark; Athena CTAS; EMR Serverless", "R4, R10, R11", "COS, SOS, EO"],
   ["**AWS Glue Data Catalog**", "Crawler de Glue; metastore Hive dentro de EMR (muere con el cluster)", "R5", "EO, COS"],
   ["**Amazon Athena**", "Redshift Serverless; un RDS analítico como en el parcial 1", "R6, R7, R9", "COS, REN, SOS"],
   ["**IAM** (`LabRole`, `LabInstanceProfile`)", "Claves de acceso guardadas en el servidor", "R7", "SEG"]],
  [4.4, 5.8, 3.4, 3.2])
P("Servicios facturables que contó el harness (R11): **Athena, EC2, EMR, Glue, Lambda y S3: 6 de 8**. Lambda no lo usa el sistema: son 5 funciones que trae el propio Learner Lab (`RedshiftOverwatch`, `RoleCreationFunction`, etc.), y el harness las cuenta igual. Se declara aquí para que se pueda verificar manualmente.", 10.5)

# ================================================================== 4. COSTOS
H("4. Proyección de costos")
P("El archivo `cost_projection.yaml` está en la raíz del repositorio; el harness suma sus partidas. **Supuestos:** operación 24×7; ~2.000 búsquedas y ~150 reservas por día; ~30 consultas diarias al panel; histórico de 2 M reservas (85 MB CSV, ~20 MB Parquet); 1 corrida de EMR al mes. Precios on-demand de us-east-1.")
T(["Servicio", "Detalle", "USD/mes"],
  [["EC2", "`t3.small` 24×7: 730 h × 0,0208", "15,18"],
   ["EC2", "Disco EBS gp3 de 8 GB (sistema + datos de PostgreSQL)", "0,64"],
   ["VPC", "Elastic IP asociada: 730 h × 0,005", "3,65"],
   ["S3", "~0,2 GB; los resultados de Athena expiran a los 7 días y los logs de EMR a los 14", "0,05"],
   ["Athena", "~1.000 consultas/mes × 10 MB mínimos ≈ 10 GB", "0,05"],
   ["EMR", "1 corrida/mes: 2 × m4.large, 1 h facturada cada uno (EC2 0,10 + EMR 0,03)", "0,26"],
   ["Glue", "Data Catalog con 2 tablas (gratuito por debajo del millón de objetos); sin crawlers", "0,00"],
   ["**Total**", "", "**19,83**"]],
  [2.4, 12.0, 2.4])
P("**Verificación del harness (R12):** el total declarado de **US$ 19,83** está por debajo del límite de US$ 50, que es el mínimo entre 1,5× la referencia de US$ 60 y el presupuesto. El costo que el harness estima a partir de los recursos vivos es **US$ 15,90**: EC2 15,18 + EMR 0,72 por las 24 horas normalizadas ya consumidas. La diferencia (US$ 3,9) está dentro de la tolerancia de 25 % + US$ 2. Se explica porque el estimador del harness no cuenta la IP elástica ni el disco EBS, que sí se pagan.")
H("4.1 ¿Qué pasa si los datos o las consultas se multiplican por 10?", 2)
T(["Escenario", "Qué cambia", "USD/mes", "Variación"],
  [["Base", "—", "19,83", "—"],
   ["Datos ×10 (20 M reservas, 850 MB CSV, ~200 MB Parquet)", "S3 +0,02; EMR con 2 nodos core (+0,13); la tabla de métricas no crece (sigue siendo rutas × meses)", "≈ 19,98", "+0,8 %"],
   ["Consultas ×10 (~9.000 al panel por mes)", "Athena: 90 GB facturados → 0,44", "≈ 20,22", "+2,0 %"],
   ["Reservas ×10 (~1.500 por día)", "Ninguno: R1 se probó con 100 peticiones a concurrencia 20 y p95 de 0,185 s", "≈ 19,83", "0 %"],
   ["Todo ×10", "Suma de lo anterior", "≈ 20,4", "+3 %"]],
  [5.0, 7.6, 2.0, 2.2])
NOTE("**Conclusión de costos y sostenibilidad:** el 98 % del gasto es capacidad **siempre encendida** (la EC2 y su IP), no consumo. El camino analítico es efímero (EMR vive ~10 minutos al mes) o serverless (Athena no cobra estando inactivo), y por eso el costo crece de forma **sublineal** con datos y consultas. La palanca que queda es la EC2: un Savings Plan de 1 año (~−30 %) o una `t4g.small` ARM (−19 %, más eficiente por vatio). La que sí se aplica en el laboratorio es apagar la EC2 cuando no se trabaja. **Pilares:** COS, SOS.")

# ================================================================== 5. IMPLEMENTACIÓN / TRAZABILIDAD
H("5. Implementación y trazabilidad")
H("5.1 Requisito → componente → evidencia", 2)
T(["Req", "Componente y archivo", "Evidencia del harness"],
  [["R1", "Índice `ix_flight_search`; búsqueda sin bloqueos · `backend/app/booking.py`", "p50 0,091 s · p95 0,185 s"],
   ["R2", "`SELECT … FOR UPDATE` + CHECK `ck_no_oversell`; pool de 10 conexiones por proceso · `booking.py`, `sql/schema.sql`", "5 rondas 1 silla: 1×201 + 49×409; 5 rondas K=8: 8×201"],
   ["R3", "`GET /reservations/{id}` · `backend/app/main.py`", "200 con la misma forma; 404 si no existe"],
   ["R4", "`emr/curate.py`, `infra/02_emr.sh`", "24/24 Parquet escritos durante una corrida de EMR; 2.005.463 filas"],
   ["R5", "`athena/ddl.sql`, `infra/03_athena.sh`", "Columnas exactas, Parquet, projection por mes"],
   ["R6", "Tabla de métricas + particiones", "Top-10 igual al oráculo; 295 KB escaneados"],
   ["R7", "`backend/app/analytics.py`", "5 sondas correctas en 1,6–2,2 s; 7 consultas en el historial de Athena"],
   ["R8", "`backend/app/static/index.html`", "Reserva creada desde Chrome y leída en el backend"],
   ["R9", "`backend/app/static/manager.html`", "LET-PSO 2024-03: panel = endpoint = oráculo = 0,7757"],
   ["R10", "`--auto-terminate` en `infra/02_emr.sh`", "0 clusters activos; 24 de 150 horas normalizadas"],
   ["R11", "t3.small, m4.large, sin NAT/ALB/Multi-AZ", "6 de 8 servicios"],
   ["R12", "`cost_projection.yaml`", "19,83 declarados frente a 15,90 en vivo"]],
  [1.2, 8.8, 6.8])
H("5.2 Control de sobreventa (R2): la transacción", 2)
P("Es el DEC-4 del primer parcial reducido a un solo tramo:")
N(["`BEGIN` y `SET LOCAL lock_timeout = 10 s`.",
   "`SELECT capacity, seats_sold … WHERE flight_id = $1 FOR UPDATE`: bloquea la fila del vuelo y las demás solicitudes hacen fila.",
   "Si `seats_sold >= capacity` → **409**. Si el vuelo no existe → **404**.",
   "`UPDATE flight SET seats_sold = seats_sold + 1` e `INSERT INTO reservation`.",
   "`COMMIT`: libera el bloqueo y la siguiente solicitud lee el contador ya actualizado."])
P("Si se agotara el tiempo de espera del bloqueo, la respuesta es 409 y no 500. El contrato exige que \"sin sillas\" nunca sea un error del servidor. El tiempo de espera es holgado a propósito: si venciera, una solicitud que sí tenía silla recibiría 409 y la prueba de K sillas contaría menos de K.")
H("5.3 Validación del job antes de gastar en el dataset completo", 2)
P("El job se corrió primero sobre una **muestra de 150 mil filas** generada con el mismo seed. `infra/validate_sample.py` descarga los Parquet y los compara con el oráculo del harness, que calcula la respuesta correcta para ese mismo tamaño. Revisa columnas, filas totales y por mes, ingresos confirmados por mes, y capacidad, sillas vendidas, ingresos, cancelaciones y ocupación de las 336 métricas. **Todo cuadró** antes de lanzar los 2 M de filas, que también cuadraron a la primera.")

# ================================================================== 6. CAMBIOS EN VIVO
H("6. Impacto de cambios (preparación para la sustentación)")
T(["Si el docente pide…", "Qué se modifica", "Qué se vuelve a correr", "Requisitos afectados"],
  [["Otra regla de limpieza (p. ej. descartar precios > 1.000.000)", "Un `filter` en `clean_reservations` de `curate.py`", "`02_emr.sh full` (~9 min). Las tablas no cambian", "R4, R6, R7, R9 (el oráculo del harness debe aplicar la misma regla)"],
   ["Otra agregación (p. ej. ingreso promedio por ruta y mes)", "Una columna nueva en `build_metrics` y en el DDL de `route_month_metrics`", "EMR + `03_athena.sh`", "R5 (las columnas nuevas van al final), R6"],
   ["Particionar por ruta además de mes", "`partitionBy('month','route')` y el DDL", "EMR + DDL", "R4 (rutas de S3), R5, R6"],
   ["Mostrar también ingresos en el panel", "`analytics.py` (SELECT) y `manager.html`", "`05_redeploy.sh` (sin EMR)", "R7 y R9 siguen pasando si se conserva el formato"],
   ["Permitir cancelar reservas", "Endpoint nuevo: `seats_sold - 1` con el mismo bloqueo", "`05_redeploy.sh`", "R2 (el invariante debe mantenerse)"]],
  [4.2, 4.6, 3.6, 4.4], 9)

# ================================================================== 7. HARNESS (SECCIÓN 6)
H("7. Ejecución del harness (sección 6)")
H("7.1 Cómo se ejecuta", 2)
N(["Copiar `config.yaml` y `cost_projection.yaml` del repositorio a la carpeta del harness (seed 190, URLs `http://100.60.245.136:8000`, base `airline_analytics`).",
   "Copiar las credenciales del Learner Lab (AWS Details → AWS CLI) en `~/.aws/credentials`. Vencen cada 4 horas y nunca se suben al repositorio.",
   "Asegurarse de que **no haya un cluster EMR vivo**. Si lo hay, sus nodos m4.large aparecen como EC2 y R11 falla.",
   "`pip install -r requirements.txt && playwright install chromium`, luego `pytest`. Se generan `report.json` y `report.md`."])
H("7.2 Resultado", 2)
T(["Req", "Título", "Estado", "Puntos"],
  [["R1", "Latencia de búsqueda", "passed", "4/4"], ["R2", "Sin sobreventa bajo concurrencia", "passed", "10/10"],
   ["R3", "Consulta de reserva", "passed", "3/3"], ["R4", "EMR produce Parquet curado", "passed", "8/8"],
   ["R5", "Tablas del Glue Catalog según contrato", "passed", "6/6"], ["R6", "Athena: corrección y poda de particiones", "passed", "8/8"],
   ["R7", "Endpoint analítico respaldado por Athena", "passed", "8/8"], ["R8", "Flujo del cliente en navegador", "passed", "6/6"],
   ["R9", "El panel coincide con el endpoint", "passed", "4/4"], ["R10", "Sin EMR remanente / horas acotadas", "passed", "5/5"],
   ["R11", "Arquitectura simple y de bajo costo", "passed", "4/4"], ["R12", "Proyección de costos", "passed", "4/4"],
   ["", "**Total**", "", "**70/70**"]],
  [1.4, 9.2, 2.8, 2.4])
P("**Lo que el harness no pudo verificar.** R11 advierte `could not inspect: ['CloudFront']`: el rol del laboratorio no puede listar distribuciones de CloudFront. El sistema no usa CloudFront y se puede comprobar manualmente en la consola. Lambda aparece contado por las funciones propias del laboratorio (sección 3.6). No hubo requisitos en estado *not verified*.")
IMG(FIG / "cliente-confirmacion.png", 11.5, "Figura 2. Flujo del cliente en el sistema desplegado: búsqueda, reserva y código de confirmación (R8).")
IMG(FIG / "panel-gerencial.png", 11.5, "Figura 3. Panel gerencial: ocupación de SMR-CTG en enero de 2024 leída de Athena (R9).")

H("7.3 Hallazgos del despliegue", 2)
P("Problemas que no se ven leyendo el diseño y que solo aparecieron al desplegar:", 10.5)
T(["#", "Síntoma", "Causa raíz", "Corrección"],
  [["H-1", "EMR de 1 nodo: el step quedó 15 min en RUNNING sin escribir nada", "Driver en modo cluster (~2,4 GB) + executor (~5 GB) no caben en los ~6 GB de YARN de un m4.large. Log: `Initial job has not accepted any resources`", "Memoria explícita: `spark.executor.memory=2g` y driver en modo client"],
   ["H-2", "Segundo intento con 1 nodo: 25 min sin avanzar y sin subir logs a S3", "Un solo m4.large aloja HDFS, YARN, los agentes de EMR, el driver y los executors: se satura", "1 master + 1 core m4.large: el job tarda 2 min"],
   ["H-3", "`syntax error near unexpected token` al final de un script de EMR", "Bash lee el script a medida que lo ejecuta; editarlo mientras corre corrompe la lectura", "Cada corrida usa una copia del script"],
   ["H-4", "El AWS CLI no encontraba las credenciales", "El Bloc de notas guardó el archivo como `credentials.txt`", "Renombrar a `credentials` (sin extensión)"],
   ["H-5", "Falló la instalación de Playwright", "Windows no admite rutas largas y la carpeta del proyecto es muy profunda", "Entorno virtual en una ruta corta (`C:\\Users\\…\\venvs\\bigdata`)"],
   ["H-6", "Riesgo: el docente califica con una URL que deja de existir", "El Learner Lab detiene la EC2 al cerrar la sesión y cambia su IP pública", "Elastic IP (sección 3.1)"],
   ["H-7", "Discrepancia en el material: `bookings`", "El oráculo cuenta confirmadas + canceladas; el contrato A.2 dice solo CONFIRMED", "Se sigue el contrato. Ningún test lo evalúa"]],
  [1.1, 4.3, 6.2, 5.2], 9)
P("Sobre H-1 y H-2: los dos intentos fallidos costaron 8 horas normalizadas (~US$ 0,30) y se detectaron sobre la muestra. Lanzar directamente los 2 M de filas habría dejado el mismo cluster atascado, con el mismo costo, pero sin un diagnóstico rápido.", 10.5)

# ================================================================== 8. DESVIACIONES
H("8. Desviaciones y riesgos declarados")
T(["#", "Desviación", "Motivo", "En producción"],
  [["D-1", "PostgreSQL sin backups automáticos", "Decisión de costo (sección 3.1)", "`pg_dump` diario a S3 por cron (~US$ 0,01/mes) o migrar a RDS"],
   ["D-2", "Una sola zona de disponibilidad", "R11 prohíbe Multi-AZ en este ejercicio", "RDS Multi-AZ o réplica en otra zona"],
   ["D-3", "Rol compartido `LabRole`, sin mínimo privilegio", "El laboratorio no permite crear roles", "Un rol para la EC2 con `athena:StartQueryExecution`/`GetQuery*` y lectura de S3 en dos prefijos"],
   ["D-4", "HTTP sin TLS en el puerto 8000, abierto a internet", "El harness del docente debe alcanzarlo; un certificado exige ALB o CloudFront (costo y servicios)", "ALB con certificado de ACM o CloudFront delante"],
   ["D-5", "El endpoint `/admin/flights` no tiene autenticación", "Lo exige el contrato para que el harness cree vuelos", "Autenticación o red privada"]],
  [1.1, 5.0, 5.2, 5.5], 9)

# ================================================================== 9. REPRODUCCIÓN
H("9. Repositorio y reproducción")
T(["Ruta", "Contenido"],
  [["`emr/curate.py`", "Job PySpark: C1–C5, cruce con vuelos, Parquet por mes y métricas"],
   ["`backend/app/`", "`main.py` (API + páginas), `booking.py` (R1–R3), `analytics.py` (R7), `static/` (frontend)"],
   ["`backend/sql/schema.sql`", "Tablas `flight` y `reservation`, CHECK anti-sobreventa, índice de búsqueda"],
   ["`athena/ddl.sql`", "Base y tablas del Glue Catalog (projection por mes)"],
   ["`infra/01_s3.sh` … `05_redeploy.sh`", "Bucket y carga, EMR, Athena, EC2 (user-data) y redespliegue por SSM"],
   ["`infra/validate_sample.py`", "Comparación de la salida de EMR con el oráculo"],
   ["`config.yaml`, `cost_projection.yaml`, `report.md`", "Configuración del harness (sin credenciales), proyección y reporte final"]],
  [5.6, 11.2])
P("Orden de ejecución desde Git Bash: `bash infra/01_s3.sh` → `bash infra/02_emr.sh sample` → `python infra/validate_sample.py s3://…/curated-sample 100000` → `bash infra/02_emr.sh full` → `bash infra/03_athena.sh` → `bash infra/04_ec2.sh`.")

# ================================================================== ANEXO A: BITÁCORA
PAGEBREAK()
H("Anexo A. Bitácora de prompts")
P("**Asistente:** Claude (Opus 5.5) en Claude Code. **Uso:** lectura del enunciado y del harness, código del backend, del frontend y del job de PySpark, scripts de AWS, diagnóstico de fallos y redacción de este documento. Cada entrada registra el prompt real, lo que produjo y **lo que el equipo decidió o corrigió**, que es la columna que importa.", 10.5)
LOG = [
    ("P-01 · Encuadre",
     "\"ayúdame a hacer mi parcial de big data, se supone que es una continuación de otro … tengo 50 dólares para gastar … primero revisemos qué hay que hacer, luego discutimos el cómo, tomamos decisiones entre ambos y luego lo hacemos\"",
     "Lectura del enunciado, del código del harness (tests, oráculo, umbrales, precios) y del parcial 1. Resumen de requisitos, siete trampas detectadas en los tests (p. ej., los RDS detenidos cuentan en R12, un cluster EMR vivo hace fallar R11) y una arquitectura propuesta.",
     "El equipo fijó el método: **primero entender, luego decidir juntos, luego construir**. No se escribió código hasta tener el seed y la decisión de base de datos."),
    ("P-02 · Decisiones de arranque",
     "\"la seed para los datos es 190 … se entrega el 23 y la base de datos que sea postgre\"",
     "Generación de los datos crudos y de las respuestas esperadas (2.005.463 filas, 336 métricas, K = 8), job de PySpark, backend, frontend y scripts.",
     "El equipo eligió PostgreSQL. El asistente propuso alojarlo en la EC2 en vez de RDS (−US$ 12/mes). Se aceptó con el compromiso de backups declarado (D-1)."),
    ("P-03 · Credenciales",
     "\"¿cómo así que las credenciales? ¿eso para qué?\"",
     "Explicación: el harness y los scripts manejan la cuenta desde el PC con credenciales temporales del laboratorio; el servidor usa el rol de la instancia y no necesita claves.",
     "El equipo pegó las credenciales en el chat; el asistente pidió no repetirlo. El archivo se guardó como `.txt` (H-4). Lección: las credenciales van solo en `~/.aws/credentials` y nunca en el repositorio ni en el chat."),
    ("P-04 · El EMR que no terminaba",
     "\"¿no se está demorando?\" — y luego: \"sucede algo, ya van 2 veces que te quedas ahí pensando\"",
     "Diagnóstico con los logs de S3 y CloudWatch: el cluster de 1 nodo no tenía memoria para los executors (H-1) y después quedó saturado (H-2). Cambio a 1 master + 1 core y monitor de estado visible.",
     "**La corrección vino del equipo:** el asistente estaba esperando en silencio. El equipo exigió visibilidad, y eso llevó a revisar los logs, a apagar el cluster atascado (que seguía gastando) y a dejar el monitor de estados. Se aprendió a no esperar a ciegas un proceso que cobra por hora."),
    ("P-05 · Entender antes de defender",
     "\"explícame qué pedía y qué hicimos para resolverlo\"",
     "Explicación requisito por requisito: qué pedía, qué se construyó, por qué y en qué estado quedó.",
     "Sirvió de base para la sustentación y para la estructura de este documento."),
    ("P-06 · IP fija",
     "\"sí, ponle la Elastic IP a la EC2 y actualiza config.yaml\"",
     "Elastic IP asociada, `config.yaml` y proyección actualizados; R8–R12 verificados de nuevo contra la IP nueva.",
     "El equipo aprobó el gasto (US$ 3,65/mes, igual que la IP automática) para que la URL de calificación no cambie (H-6)."),
    ("P-07 · Documento",
     "\"arranca con el documento Word (secciones 5 y 6, diagrama, tabla de servicios, costos y bitácora de prompts)\"",
     "Este documento, el diagrama de arquitectura y las capturas del sistema real.",
     "[[Completar: qué revisó, cambió o rechazó el equipo al leer el borrador.]]"),
]
for title, prompt, produced, correction in LOG:
    H(title, 2)
    T(["Campo", "Contenido"],
      [["Prompt", prompt], ["Qué produjo", produced], ["Decisión o corrección del equipo", correction]],
      [3.8, 13.0], 9.5)

doc.save(OUT)
print(f"ok -> {OUT}")
